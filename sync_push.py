# -*- coding: utf-8 -*-
"""把《余额.xlsx》导出成 data.json 并推送到 GitHub Pages。

自包含版本：不再依赖外部 server.py，直接在仓库目录里运行即可。
    python sync_push.py            # 导出 + 推送
    python sync_push.py --no-push  # 只导出，不推送
"""
from __future__ import annotations

import datetime as dt
import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent            # .../my-money
BRANCH = "main"
BOOK = "余额.xlsx"

# 可能的表格位置（按优先级）
CANDIDATES = [
    ROOT.parent / BOOK,                            # 仓库旁边（推荐布局）
    ROOT / BOOK,
    Path.home() / "Desktop" / BOOK,
    Path("E:/AAA-DeskBox/我的桌面/余额") / BOOK,
    Path("C:/Users/33431/Desktop") / BOOK,
]


# ---------------------------------------------------------------- 找表格
def find_book() -> Path:
    override = ROOT / "xlsx_path.txt"
    if override.exists():
        p = Path(override.read_text(encoding="utf-8").strip())
        if p.exists():
            return p
    for p in CANDIDATES:
        if p.exists():
            return p
    hits = sorted(
        (p for p in Path.home().rglob(BOOK) if p.is_file()),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if hits:
        return hits[0]
    raise SystemExit("找不到《余额.xlsx》，请把表格放到 %s 旁边" % ROOT.parent)


# ---------------------------------------------------------------- 读表格
def ensure_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("正在安装 openpyxl …")
        subprocess.run([sys.executable, "-m", "pip", "install", "-q", "openpyxl"], check=False)


def pick_sheet(wb):
    want = str(dt.date.today().year)
    if want in wb.sheetnames:
        return wb[want], want
    best, score = None, -1
    for ws in wb.worksheets:
        n = sum(1 for row in ws.iter_rows(min_row=2, max_col=1, values_only=True) if row[0])
        if n > score:
            best, score = ws, n
    if best is None:
        raise SystemExit("表格里没有可用工作表")
    return best, best.title


def date_text(v) -> str:
    if isinstance(v, (dt.datetime, dt.date)):
        return "%d.%d" % (v.month, v.day)
    return str(v).strip()


def num(v):
    if v is None or v == "":
        return None
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def build_payload(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws, sheet_name = pick_sheet(wb)

    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    header = [str(h).strip() if h is not None else "" for h in header]
    account_keys = [h for h in header[1:-2] if h]

    rows = []
    for r in range(2, ws.max_row + 1):
        raw_date = ws.cell(r, 1).value
        if raw_date is None or str(raw_date).strip() == "":
            continue
        accounts = {key: num(ws.cell(r, 2 + i).value) for i, key in enumerate(account_keys)}
        total = num(ws.cell(r, 2 + len(account_keys)).value)
        spend = num(ws.cell(r, 3 + len(account_keys)).value)

        empty = all(v is None for v in accounts.values())
        if empty:
            total = None
            spend = None
        elif total is None:
            total = round(sum(v for v in accounts.values() if v is not None), 2)

        rows.append(
            {
                "row": r,
                "date": date_text(raw_date),
                "accounts": accounts,
                "total": total,
                "empty": empty,
                "spend": spend,
            }
        )

    filled = [x for x in rows if not x["empty"]]
    if not filled:
        raise SystemExit("表格里还没有任何已填写的记录")

    latest = filled[-1]
    week = [x["spend"] for x in filled[-7:] if x["spend"] is not None]

    return {
        "headers": header,
        "account_keys": account_keys,
        "rows": rows,
        "summary": {
            "latest_date": latest["date"],
            "latest_row": latest["row"],
            "total": latest["total"],
            "accounts": latest["accounts"],
            "week_spend": round(sum(week), 2) if week else 0,
            "filled_count": len(filled),
            "row_count": len(rows),
            "sheet": sheet_name,
            "generated_at": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        },
    }


# ---------------------------------------------------------------- 导出
def export(book: Path) -> dict:
    payload = build_payload(book)
    (ROOT / "data.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    s = payload["summary"]
    print("表格：%s  [%s]" % (book, s["sheet"]))
    print("导出：data.json  最新日期 %s  全部资产 %s" % (s["latest_date"], s["total"]))
    return payload


# ---------------------------------------------------------------- 推送
def run_git(args, **kw):
    return subprocess.run(
        ["git", "-c", "user.name=Bury-HUI",
         "-c", "user.email=Bury-HUI@users.noreply.github.com"] + args,
        cwd=str(ROOT), text=True, **kw,
    )


def push(payload) -> bool:
    files = ["data.json", "sync_push.py", "watch_sync.py", "sync-push.bat", "README.md"]
    files = [f for f in files if (ROOT / f).exists()]
    run_git(["add"] + files, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    staged = run_git(["diff", "--cached", "--quiet"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    if staged.returncode == 0:
        print("数据没有变化，跳过提交。")
        return True

    msg = "update balance data %s" % payload["summary"].get("latest_date", "")
    commit = run_git(["commit", "-m", msg], capture_output=True)
    if commit.returncode != 0:
        print("提交失败：", (commit.stderr or "").strip()[:300])
        return False
    print("提交：", msg)

    current = run_git(["branch", "--show-current"], capture_output=True).stdout.strip()
    if current != BRANCH:
        run_git(["branch", "-M", BRANCH], capture_output=True)

    token_file = ROOT / "token.txt"
    token = token_file.read_text(encoding="utf-8").strip() if token_file.exists() else ""
    if token:
        url = "https://Bury-HUI:%s@github.com/Bury-HUI/my-money.git" % token
        r = run_git(["push", url, "HEAD:" + BRANCH], capture_output=True)
    else:
        r = run_git(["push", "-u", "origin", BRANCH], capture_output=True)
    if r.returncode != 0:
        err = (r.stderr or r.stdout or "").strip()
        print("推送失败：")
        print(err[:800])
        print("\n如果提示需要凭据：把 GitHub classic PAT 粘到 %s 后重跑本脚本。" % (ROOT / "token.txt"))
        return False

    print("推送成功 → https://bury-hui.github.io/my-money/")
    print("（GitHub Pages 一般 30~60 秒后生效，网页点右上角 ⟳ 可强制刷新）")
    return True


def main() -> int:
    ensure_openpyxl()
    book = find_book()
    payload = export(book)
    if "--no-push" in sys.argv:
        return 0
    return 0 if push(payload) else 1


if __name__ == "__main__":
    raise SystemExit(main())
